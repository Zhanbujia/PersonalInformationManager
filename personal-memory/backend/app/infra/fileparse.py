"""把常见文档抽成纯文本。

只做「抽文本」这一件事，不做分析——抽出来的文本仍走同一条 ingest 流水线
（规范化 → LLM 结构化分析 → 切片向量 → 建关联），不重复造一套。

解析失败一律抛异常，由流水线写成 `parse_error`，**不静默吞掉**：
用户需要知道"这份文件为什么没被整理"，而不是看到一个永远 pending 的记录。
"""

from __future__ import annotations

import html
import re
from pathlib import Path

# 允许上传的扩展名（小写，含点）。图片不在内：OCR 需要额外系统依赖（tesseract），
# 不做半吊子支持，宁可在上传时就明确告诉用户不支持。
SUPPORTED_EXTS = frozenset(
    {
        ".pdf",
        ".docx",
        ".xlsx",
        ".pptx",
        ".txt",
        ".md",
        ".markdown",
        ".csv",
        ".tsv",
        ".json",
        ".log",
        ".yaml",
        ".yml",
        ".html",
        ".htm",
        ".xml",
    }
)

_PLAIN_EXTS = frozenset({".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".yaml", ".yml"})
_MARKUP_EXTS = frozenset({".html", ".htm", ".xml"})

_SCRIPT_STYLE = re.compile(r"<(script|style)[^>]*>.*?</\1>", re.IGNORECASE | re.DOTALL)
_BLOCK_END = re.compile(r"</(p|div|br|li|tr|h[1-6]|section|article|table|ul|ol)>", re.IGNORECASE)
_TAG = re.compile(r"<[^>]+>")
_MULTI_BLANK = re.compile(r"\n{3,}")
_TRAILING_WS = re.compile(r"[ \t]+\n")


class UnsupportedFileType(Exception):
    """不在白名单里的文件类型。"""


class FileParseError(Exception):
    """文件打不开或抽不出文本（损坏、加密、扫描件等）。"""


def is_supported(ext: str) -> bool:
    return ext.lower() in SUPPORTED_EXTS


def extract_text(path: str | Path, ext: str, *, max_chars: int = 200_000) -> str:
    """按扩展名抽取文本。失败抛 FileParseError / UnsupportedFileType。"""
    ext = (ext or "").lower()
    if not ext.startswith("."):
        ext = f".{ext}"
    if ext not in SUPPORTED_EXTS:
        raise UnsupportedFileType(f"暂不支持的文件类型：{ext}")

    p = Path(path)
    if not p.exists() or not p.is_file():
        raise FileParseError("文件不存在")

    if ext == ".pdf":
        text = _extract_pdf(p)
    elif ext == ".docx":
        text = _extract_docx(p)
    elif ext == ".xlsx":
        text = _extract_xlsx(p)
    elif ext == ".pptx":
        text = _extract_pptx(p)
    elif ext in _PLAIN_EXTS:
        text = _read_text(p)
    elif ext in _MARKUP_EXTS:
        text = _strip_markup(_read_text(p))
    else:  # pragma: no cover - 白名单已覆盖
        raise UnsupportedFileType(f"暂不支持的文件类型：{ext}")

    text = _clean(text)
    if max_chars and len(text) > max_chars:
        text = text[:max_chars] + "\n\n……（内容过长，已截断）……"
    return text


# ---------------------------------------------------------------- 文本类

def _decode(data: bytes) -> str:
    """编码识别：先 charset_normalizer，再按常见中文编码兜底。"""
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        if best is not None:
            return str(best)
    except Exception:  # noqa: BLE001 - 识别失败就走后面的兜底
        pass
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _read_text(path: Path) -> str:
    try:
        return _decode(path.read_bytes())
    except OSError as exc:
        raise FileParseError(f"读取文件失败：{exc}") from exc


def _strip_markup(s: str) -> str:
    s = _SCRIPT_STYLE.sub(" ", s)
    s = _BLOCK_END.sub("\n", s)
    s = _TAG.sub(" ", s)
    s = html.unescape(s)
    return s


# ---------------------------------------------------------------- Office / PDF

def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise FileParseError("缺少 pypdf 依赖，无法解析 PDF") from exc

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise FileParseError("PDF 已加密，无法解析") from exc
        pages: list[str] = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001 - 单页坏了不影响其它页
                pages.append("")
    except FileParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise FileParseError(f"无法打开 PDF：{exc}") from exc

    text = "\n\n".join(pages)
    if not text.strip():
        raise FileParseError("PDF 未提取到文字（可能是扫描件，需要 OCR）")
    return text


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise FileParseError("缺少 python-docx 依赖，无法解析 Word") from exc

    try:
        doc = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise FileParseError(f"无法打开 Word 文档：{exc}") from exc

    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path, *, max_rows: int = 3000, max_cols: int = 60) -> str:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise FileParseError("缺少 openpyxl 依赖，无法解析 Excel") from exc

    try:
        wb = openpyxl.load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise FileParseError(f"无法打开 Excel：{exc}") from exc

    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            parts.append(f"## 工作表：{ws.title}")
            n = 0
            for row in ws.iter_rows(values_only=True):
                if n >= max_rows:
                    parts.append("……（该表行数过多，已截断）……")
                    break
                vals = ["" if v is None else str(v) for v in row[:max_cols]]
                if any(vals):
                    parts.append(" | ".join(vals).rstrip(" |"))
                    n += 1
    finally:
        wb.close()
    return "\n".join(parts)


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise FileParseError("缺少 python-pptx 依赖，无法解析 PPT") from exc

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise FileParseError(f"无法打开 PPT：{exc}") from exc

    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"## 第 {i} 页")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(r.text for r in para.runs).strip()
                    if line:
                        parts.append(line)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [(c.text or "").strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
    return "\n".join(parts)


# ---------------------------------------------------------------- 收尾

def _clean(s: str) -> str:
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = _TRAILING_WS.sub("\n", s)
    s = _MULTI_BLANK.sub("\n\n", s)
    return s.strip()
